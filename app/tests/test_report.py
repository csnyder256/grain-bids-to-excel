"""Report builder: structural round-trip, formats, zero-external-links."""
import zipfile
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook

from bidboard.engine.store import Store
from bidboard.engine.types import BidRow, ExtractionResult, FetchResult
from bidboard.report import builder, styles
from bidboard.settings import DEFAULT_SETTINGS
import copy

NOW = datetime(2026, 7, 4, 8, 30)


def _fetch():
    return FetchResult(True, "https://x.example/bids", "https://x.example/bids",
                       1, 200, "<html></html>", datetime.now(timezone.utc), 10)


def _row(commodity, delivery, start, cash, basis, futures=4.23, rh=None):
    r = BidRow(
        commodity=commodity, commodity_raw=commodity, location_label="Westfield",
        delivery_raw=delivery, delivery_label=delivery,
        delivery_start=date.fromisoformat(start),
        delivery_end=date.fromisoformat(start),
        futures_month="2026-09", cash_price=cash, basis=basis,
        futures_price=futures, change=0.0025,
    )
    from bidboard.engine import staleness as S
    r.row_hash = rh or S.row_identity_hash(1, r)
    r.value_hash = S.value_hash(r)
    return r


@pytest.fixture
def seeded_store(tmp_path):
    store = Store(str(tmp_path / "t.db"))
    c1 = store.add_company("Prairie Ridge Grain")
    c2 = store.add_company("Prairie Grain Co-op")
    s1 = store.add_site(c1, "https://prairieridge.example/2163", "Westfield")
    s2 = store.add_site(c2, "https://prairie.example/bids", "Main")
    run = store.begin_run()
    ex1 = ExtractionResult(rows=[
        _row("CORN", "July 1 - 15", "2026-07-01", 4.10, -0.13, rh="a1"),
        _row("CORN", "August 1 - 15", "2026-08-01", 4.14, -0.09, rh="a2"),
        _row("SOYBEANS", "October 2026", "2026-10-01", 10.25, -0.45, rh="a3"),
    ], as_of_date=date(2026, 7, 2))
    ps1 = store.record_page_scan(run, s1, _fetch(), ex1, "ok", content_signature="sig1")
    store.save_bids(ps1, s1, c1, ex1.rows)
    ex2 = ExtractionResult(rows=[
        _row("CORN", "July 1 - 15", "2026-07-01", 4.08, -0.15, rh="b1"),
    ], as_of_date=date(2026, 7, 1))
    ps2 = store.record_page_scan(run, s2, _fetch(), ex2, "ok", content_signature="sig2")
    store.save_bids(ps2, s2, c2, ex2.rows)
    store.finish_run(run, "complete")
    yield store, run
    store.close()


def test_master_workbook_structure(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)

    assert len(result.files) == 1
    path = result.files[0].path
    wb = load_workbook(path)

    assert "Summary" in wb.sheetnames
    assert "Data Health" in wb.sheetnames
    assert "Prairie Ridge Grain" in wb.sheetnames
    assert "All Bids" in wb.sheetnames

    cover = wb["Summary"]
    assert cover["A1"].value == "Cash Bid Report"

    bidsheet = wb["Prairie Ridge Grain"]
    assert bidsheet["A1"].value == "Prairie Ridge Grain"
    assert bidsheet.freeze_panes == "A7"
    assert bidsheet.auto_filter.ref is not None
    assert bidsheet.sheet_view.showGridLines is False


def test_number_formats_and_values(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    wb = load_workbook(result.files[0].path)
    ws = wb["Prairie Ridge Grain"]

    header_row = 6
    headers = [ws.cell(row=header_row, column=j).value for j in range(1, 12)]
    assert "Cash Bid" in headers
    assert "Basis" in headers

    cash_col = headers.index("Cash Bid") + 1
    basis_col = headers.index("Basis") + 1
    cash_cell = ws.cell(row=header_row + 1, column=cash_col)
    basis_cell = ws.cell(row=header_row + 1, column=basis_col)
    assert cash_cell.number_format == styles.fmt_cash(4)
    assert basis_cell.number_format == styles.fmt_basis(4)
    # values present
    cash_values = [ws.cell(row=r, column=cash_col).value
                   for r in range(header_row + 1, header_row + 6)]
    assert 4.10 in cash_values


def test_no_external_links_no_formulas(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    settings["output"]["charts"] = "both"  # exercise chart path too
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    path = result.files[0].path

    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert not any(n.startswith("xl/externalLinks/") for n in names)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert not cell.value.startswith("="), f"formula in {ws.title}!{cell.coordinate}"


def test_charts_created_when_enabled(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    settings["output"]["charts"] = "bars"
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    wb = load_workbook(result.files[0].path)
    assert "_ChartData" in wb.sheetnames
    assert wb["_ChartData"].sheet_state == "hidden"


def test_per_company_mode_one_file_each(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    settings["output"]["workbook_modes"] = ["per_company"]
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    assert len(result.files) == 2
    scopes = {f.scope for f in result.files}
    assert "Prairie Ridge Grain" in scopes
    assert "Prairie Grain Co-op" in scopes


def test_never_overwrites(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    out = tmp_path / "out"
    r1 = builder.build_reports(store, settings, run, out, now=NOW)
    r2 = builder.build_reports(store, settings, run, out, now=NOW)
    assert r1.files[0].path != r2.files[0].path  # second got a " (2)" name


# ---- naming a location the app never labelled -------------------------------
#
# Every sheet that names a location reads ``sites.label``: the bid sheets'
# Location column, the cover's "Freshness at a glance" and "Best cash bids",
# and the Data Health sheet. But the app's own add-company flow posts bare page
# URLs - ``companies.js`` sends ``sites: [{url}]`` for each page it found, and
# nothing in the UI ever calls the one endpoint that accepts a label. So
# ``sites.label`` is NULL for every location a user adds, and the Location
# column came out blank on every row of a real workbook: two elevators' bids
# sitting side by side with nothing saying which was which. A location that has
# no name is still one specific page, so it is now named by that page's address
# - the same spelling the Data Health sheet already uses in its "Page" column.

NAMELESS_URL = "https://prairieridge.example/cashbidssingle-2164"
NAMELESS_DISPLAY = "prairieridge.example/cashbidssingle-2164"


def _col(ws, header, header_row=6):
    """Index of the column whose header cell is ``header``."""
    for j in range(1, ws.max_column + 1):
        if ws.cell(row=header_row, column=j).value == header:
            return j
    raise AssertionError(f"no {header!r} column on {ws.title!r}")


def _scan(store, run, cid, sid, cash):
    # imported here, like ``_row`` above, so this file's import block stays as
    # the rest of the suite has it
    from bidboard.engine import staleness as S

    rows = [_row("CORN", "July 1 - 15", "2026-07-01", cash, -0.15)]
    S.stamp_hashes(sid, rows)
    ex = ExtractionResult(rows=rows, as_of_date=date(2026, 7, 2))
    ps = store.record_page_scan(run, sid, _fetch(), ex, "ok",
                                content_signature=S.content_signature(rows))
    store.save_bids(ps, sid, cid, rows)
    return rows


@pytest.fixture
def mixed_label_store(tmp_path):
    """One location named "Westfield", one added the way the UI adds them: a
    bare page URL, no label. The nameless page holds the better bid."""
    store = Store(str(tmp_path / "mixed.db"))
    cid = store.add_company("Prairie Ridge Grain")
    named = store.add_site(cid, "https://prairieridge.example/cashbidssingle-2163",
                           "Westfield")
    nameless = store.add_site(cid, NAMELESS_URL)
    run = store.begin_run()
    _scan(store, run, cid, named, 3.92)
    _scan(store, run, cid, nameless, 4.08)
    store.finish_run(run, "complete")
    yield store, run
    store.close()


def test_location_display_falls_back_to_the_page_address():
    """A label wins; with no label the page's address is the name."""
    from bidboard.report.naming import location_display, page_address

    assert location_display("Westfield", "https://x.example/bids") == "Westfield"
    assert location_display(None, NAMELESS_URL) == NAMELESS_DISPLAY
    assert location_display("   ", NAMELESS_URL) == NAMELESS_DISPLAY
    assert page_address("https://x.example/bids/") == "x.example/bids"
    assert page_address(None) == ""


def test_bid_sheet_names_an_unlabelled_location_by_its_page(mixed_label_store, tmp_path):
    """Two elevators on one company sheet must be told apart.

    Regression: both rows carried a blank Location, so the workbook showed two
    identical-looking corn bids with different prices and no way to know which
    elevator each came from.
    """
    store, run = mixed_label_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    ws = load_workbook(result.files[0].path)["Prairie Ridge Grain"]
    loc, cash = _col(ws, "Location"), _col(ws, "Cash Bid")

    named_by_cash = {}
    for r in range(7, ws.max_row + 1):
        value = ws.cell(row=r, column=cash).value
        if isinstance(value, (int, float)):
            named_by_cash[value] = ws.cell(row=r, column=loc).value

    assert named_by_cash[3.92] == "Westfield", "a real label still wins"
    assert named_by_cash[4.08] == NAMELESS_DISPLAY, "not blank"


def test_cover_and_data_health_name_an_unlabelled_location(mixed_label_store, tmp_path):
    """The cover's freshness and best-bid blocks, and Data Health, all name it.

    The cover's "Freshness at a glance" block has no page column at all, so a
    blank location there left the row completely anonymous.
    """
    store, run = mixed_label_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    wb = load_workbook(result.files[0].path)

    cover = wb["Summary"]
    locations = [cover.cell(row=r, column=2).value
                 for r in range(1, cover.max_row + 1)]
    assert NAMELESS_DISPLAY in locations
    assert "Westfield" in locations

    best_header = next(r for r in range(1, cover.max_row + 1)
                       if cover.cell(row=r, column=1).value == "Commodity")
    best_row = best_header + 1
    assert cover.cell(row=best_row, column=1).value == "CORN"
    assert cover.cell(row=best_row, column=7).value == NAMELESS_DISPLAY

    dh = wb["Data Health"]
    loc, page = _col(dh, "Location"), _col(dh, "Page")
    named_by_page = {}
    for r in range(7, dh.max_row + 1):
        key = dh.cell(row=r, column=page).value
        if key:
            named_by_page[key] = dh.cell(row=r, column=loc).value
    assert named_by_page[NAMELESS_DISPLAY] == NAMELESS_DISPLAY
    # a labelled location keeps its own name, distinct from its page address
    assert named_by_page["prairieridge.example/cashbidssingle-2163"] == "Westfield"


def test_cover_summary_counts_one_company_and_one_bid(tmp_path):
    """"1 company, 1 bid" - the cover counted them as "1 companies, 1 bids"."""
    store = Store(str(tmp_path / "one.db"))
    cid = store.add_company("Solo Grain")
    sid = store.add_site(cid, "https://solo.example/bids", "Main")
    run = store.begin_run()
    _scan(store, run, cid, sid, 4.00)
    store.finish_run(run, "complete")
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    subtitle = load_workbook(result.files[0].path)["Summary"]["A2"].value
    store.close()
    assert "1 company, 1 bid" in subtitle


def test_cover_summary_keeps_plurals_for_several(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    subtitle = load_workbook(result.files[0].path)["Summary"]["A2"].value
    assert "2 companies, 4 bids" in subtitle


def test_company_override_applies(seeded_store, tmp_path):
    store, run = seeded_store
    settings = copy.deepcopy(DEFAULT_SETTINGS)
    # restrict Prairie Ridge (company id 1) to SOYBEANS only
    settings["company_overrides"] = {"1": {"commodities": ["SOYBEANS"]}}
    result = builder.build_reports(store, settings, run, tmp_path / "out", now=NOW)
    wb = load_workbook(result.files[0].path)
    ws = wb["Prairie Ridge Grain"]
    # only soybean rows -> commodity column should show SOYBEANS, not CORN
    commodities = set()
    for r in range(7, 20):
        v = ws.cell(row=r, column=1).value
        if v:
            commodities.add(v)
    assert "CORN" not in commodities

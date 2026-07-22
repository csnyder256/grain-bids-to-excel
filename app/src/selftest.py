"""Post-install self-test. Run by SETUP.bat after dependencies install.
Prints one OK line per check; on success writes data/.setup-complete.
Any failure exits non-zero so SETUP shows its friendly guidance."""
from __future__ import annotations

import socket
import sys
import tempfile
from pathlib import Path

SRC = Path(__file__).resolve().parent
sys.path.insert(0, str(SRC))


def check(label, fn):
    try:
        fn()
        print(f"OK   {label}")
        return True
    except Exception as e:
        print(f"FAIL {label}: {e}")
        return False


def _imports():
    import apscheduler  # noqa
    import bs4  # noqa
    import flask  # noqa
    import jsonschema  # noqa
    import lxml  # noqa
    import openpyxl  # noqa
    import rapidfuzz  # noqa
    import requests  # noqa
    import waitress  # noqa


def _db_migrations():
    from bidboard.engine.store import Store
    s = Store(":memory:")
    assert s.get_meta("schema_version") == "1"
    s.close()


def _xlsx_roundtrip():
    from openpyxl import Workbook, load_workbook
    tmp = Path(tempfile.gettempdir()) / "bidboard_selftest.xlsx"
    wb = Workbook()
    wb.active["A1"] = "ok"
    wb.save(tmp)
    assert load_workbook(tmp)["Sheet"]["A1"].value == "ok"
    tmp.unlink(missing_ok=True)


def _port_bind():
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))


def _app_builds():
    from bidboard import create_app
    app = create_app(db_path=":memory:",
                     settings_path=Path(tempfile.gettempdir()) / "bidboard_selftest.json")
    with app.test_client() as c:
        assert c.get("/api/ping").status_code == 200


def main() -> int:
    checks = [
        ("Python components installed", _imports),
        ("Database ready", _db_migrations),
        ("Excel writer working", _xlsx_roundtrip),
        ("Network port available", _port_bind),
        ("App starts and answers", _app_builds),
    ]
    ok = all(check(label, fn) for label, fn in checks)

    # Playwright is optional - report but never fail on it
    try:
        import importlib
        importlib.import_module("playwright.sync_api")
        print("OK   Mini-browser (for JavaScript sites) available")
    except ImportError:
        print("NOTE Mini-browser not installed yet - some sites may need it "
              "(install later from Settings).")

    if ok:
        from bidboard import paths
        paths.ensure_runtime_dirs()
        paths.SETUP_MARKER.write_text("ok", encoding="utf-8")
        print("\nAll checks passed.")
        return 0
    print("\nOne or more checks failed.")
    return 1


if __name__ == "__main__":
    sys.exit(main())

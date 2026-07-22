"""Sheet writers - title blocks, bid tables, cover, data-health, all-bids.
Every style comes from styles.py - constants for the fixed theme, and
styles.badge(status) for the data-dependent status badges. This module
never constructs an openpyxl style object of its own."""
from __future__ import annotations

from datetime import date, datetime

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from . import styles as ST
from . import charts as CH

# canonical column registry: key -> (header, width, kind)
COLUMNS = {
    "company": ("Company", 22, "text"),
    "location": ("Location", 16, "text"),
    "commodity": ("Commodity", 15, "text"),
    "delivery": ("Delivery", 20, "text"),
    "start": ("Starts", 11, "date"),
    "end": ("Ends", 11, "date"),
    "cash": ("Cash Bid", 12, "cash"),
    "basis": ("Basis", 10, "basis"),
    "fut_month": ("Futures Month", 14, "text"),
    "futures": ("Futures", 11, "futures"),
    "change": ("Change", 10, "change"),
    "asof": ("Bids As Of", 12, "date"),
    "notes": ("Notes", 34, "wrap"),
}


def _fmt_for(kind: str, dec: dict) -> str | None:
    return {
        "cash": ST.fmt_cash(dec["cash"]),
        "basis": ST.fmt_basis(dec["basis"]),
        "futures": ST.fmt_futures(dec["futures"]),
        "change": ST.fmt_change(dec["basis"]),
        "date": ST.FMT_DATE,
    }.get(kind)


def _parse_date(value):
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def column_keys(eff: dict, include_company: bool, include_location: bool) -> list[str]:
    cols = []
    if include_company:
        cols.append("company")
    if include_location:
        cols.append("location")
    cols.append("commodity")
    cols.append("delivery")
    c = eff.get("columns", {})
    if c.get("include_delivery_dates", True):
        cols += ["start", "end"]
    cols.append("cash")
    cols.append("basis")
    if c.get("include_futures", True):
        cols += ["fut_month", "futures"]
    if c.get("include_change", True):
        cols.append("change")
    if include_company or include_location:
        cols.append("asof")
    if c.get("include_flags", True):
        cols.append("notes")
    return cols


def _title_block(ws, ncols, title, subtitle, small, source_url, eff):
    last_col = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = title
    c.font = ST.F_TITLE
    c.alignment = ST.A_TITLE
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value = subtitle
    c.font = ST.F_SUBTITLE
    ws.row_dimensions[2].height = 16

    ws.merge_cells(f"A3:{last_col}3")
    c = ws["A3"]
    c.value = small
    c.font = ST.F_SMALL
    if source_url and eff.get("hyperlinks", True):
        c.hyperlink = Hyperlink(ref="A3", target=source_url)
    ws.row_dimensions[3].height = 14

    # gold divider rule
    for col in range(1, ncols + 1):
        ws.cell(row=4, column=col).fill = ST.FILL_RULE
    ws.row_dimensions[4].height = 3


def write_bid_sheet(ws, rows, eff, *, title, subtitle, small, source_url,
                    include_company, include_location, charts_mode="none",
                    now=None):
    now = now or datetime.now()
    dec = eff.get("decimals", {"cash": 4, "basis": 4, "futures": 4})
    keys = column_keys(eff, include_company, include_location)
    ncols = len(keys)
    ws.sheet_view.showGridLines = False

    _title_block(ws, ncols, title, subtitle, small, source_url, eff)

    header_row = 6
    for j, key in enumerate(keys, start=1):
        head, width, _ = COLUMNS[key]
        cell = ws.cell(row=header_row, column=j, value=head)
        cell.font = ST.F_HEADER
        cell.fill = ST.FILL_HEADER
        cell.border = ST.B_HEADER
        _, _, kind = COLUMNS[key]
        cell.alignment = ST.A_RIGHT if kind in ("cash", "basis", "futures", "change") else ST.A_LEFT
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[header_row].height = 20

    group_by = eff.get("group_by", "commodity")
    prev_group = None
    r = header_row + 1
    cash_col_idx = keys.index("cash") + 1 if "cash" in keys else None
    basis_col_idx = keys.index("basis") + 1 if "basis" in keys else None
    change_col_idx = keys.index("change") + 1 if "change" in keys else None
    data_start = r

    for row in rows:
        group_val = _group_value(row, group_by)
        is_group_break = group_by != "none" and group_val != prev_group
        prev_group = group_val
        band = ST.FILL_BAND if ((r - data_start) % 2 == 1) else None
        for j, key in enumerate(keys, start=1):
            _, _, kind = COLUMNS[key]
            value = _cell_value(row, key)
            cell = ws.cell(row=r, column=j, value=value)
            cell.font = ST.F_BODY
            fmt = _fmt_for(kind, dec)
            if fmt:
                cell.number_format = fmt
            cell.alignment = (ST.A_RIGHT if kind in ("cash", "basis", "futures", "change")
                              else ST.A_WRAP if kind == "wrap" else ST.A_LEFT)
            cell.border = ST.B_GROUPTOP if is_group_break else ST.B_ROW
            if band and not is_group_break:
                cell.fill = band
        r += 1
    data_end = r - 1

    if data_end >= data_start:
        last_col = get_column_letter(ncols)
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{last_col}{data_end}"
        # conditional formatting
        if basis_col_idx:
            col = get_column_letter(basis_col_idx)
            rng = f"{col}{data_start}:{col}{data_end}"
            for rule in ST.cf_pos_neg(rng):
                ws.conditional_formatting.add(rng, rule)
        if change_col_idx:
            col = get_column_letter(change_col_idx)
            rng = f"{col}{data_start}:{col}{data_end}"
            for rule in ST.cf_pos_neg(rng):
                ws.conditional_formatting.add(rng, rule)
        if cash_col_idx and eff.get("columns", {}).get("highlight_best_bids", True):
            col = get_column_letter(cash_col_idx)
            rng = f"{col}{data_start}:{col}{data_end}"
            ws.conditional_formatting.add(rng, ST.cf_cash_scale())

    # footer
    foot_row = data_end + 2
    ws.merge_cells(f"A{foot_row}:{get_column_letter(ncols)}{foot_row}")
    fc = ws.cell(row=foot_row, column=1)
    domain = (source_url or "").split("//")[-1].split("/")[0]
    fc.value = (f"Source: {domain} · Generated {now:%m/%d/%Y %I:%M %p} · "
                "Bids are indicative - confirm with the elevator before selling.")
    fc.font = ST.F_SMALL

    # print setup
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"

    # charts
    if charts_mode != "none" and data_end >= data_start:
        _add_sheet_charts(ws, rows, charts_mode)


def _add_sheet_charts(ws, rows, mode):
    wb = ws.parent
    data_ws = CH.ensure_chart_data_sheet(wb)
    # group by commodity
    by_commodity: dict[str, list[dict]] = {}
    for row in rows:
        by_commodity.setdefault(row.get("commodity") or "?", []).append(row)
    anchor_row = 6
    anchor_col = ws.max_column + 2
    made = 0
    for commodity, crows in by_commodity.items():
        crows = [r for r in crows if r.get("cash_price") is not None]
        if len(crows) < 2:
            continue
        cats = [r.get("delivery_label") or "" for r in crows]
        anchor = f"{get_column_letter(anchor_col)}{anchor_row + made * 16}"
        if mode in ("bars", "both"):
            if CH.add_bid_bar_chart(wb, ws, data_ws, f"{commodity} - cash bid by delivery",
                                    cats, [r["cash_price"] for r in crows], anchor):
                made += 1
        if mode in ("lines", "both"):
            basis_vals = [r.get("basis") or 0 for r in crows]
            anchor2 = f"{get_column_letter(anchor_col)}{anchor_row + made * 16}"
            if CH.add_basis_line_chart(wb, ws, data_ws, f"{commodity} - basis across deliveries",
                                       cats, basis_vals, anchor2):
                made += 1


def _group_value(row, group_by):
    return {
        "commodity": row.get("commodity"),
        "location": row.get("site_label"),
        "company": row.get("company_name"),
    }.get(group_by, None)


def _cell_value(row, key):
    if key == "company":
        return row.get("company_name")
    if key == "location":
        return row.get("site_label")
    if key == "commodity":
        return row.get("commodity") or row.get("commodity_raw")
    if key == "delivery":
        return row.get("delivery_label") or row.get("delivery_raw")
    if key == "start":
        return _parse_date(row.get("delivery_start"))
    if key == "end":
        return _parse_date(row.get("delivery_end"))
    if key == "cash":
        return row.get("cash_price")
    if key == "basis":
        return row.get("basis")
    if key == "fut_month":
        return row.get("futures_month")
    if key == "futures":
        return row.get("futures_price")
    if key == "change":
        return row.get("change")
    if key == "asof":
        return _parse_date(row.get("page_as_of"))
    if key == "notes":
        return _notes_text(row)
    return None


def _notes_text(row):
    import json
    raw = row.get("notes_json")
    if not raw:
        return None
    try:
        notes = json.loads(raw)
    except (ValueError, TypeError):
        return None
    friendly = {
        "year_inferred": "delivery year assumed from page date",
        "unit_assumed_cents": "converted from cents",
        "newcrop_window_assumed": "new-crop window estimated",
        "commodity_unrecognized": "commodity name not recognized",
        "delivery_unparsed": "delivery period couldn't be read",
        "basis_out_of_range": "basis value looked off",
    }
    msgs = [friendly.get(n) for n in notes if friendly.get(n)]
    return "; ".join(msgs) if msgs else None


def write_cover(ws, meta: dict, health: list, best: list, needs_attention: list, eff, now):
    ws.sheet_view.showGridLines = False
    ncols = 8
    for i, w in enumerate([22, 16, 13, 9, 11, 8, 16, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    _title_block(
        ws, ncols, "Cash Bid Report",
        f"Generated {now:%A, %B %d, %Y %I:%M %p} · {meta['company_count']} companies, {meta['bid_count']} bids",
        "",
        None, eff,
    )

    r = 6
    ws.cell(row=r, column=1, value="Freshness at a glance").font = ST.F_H2
    r += 1
    heads = ["Company", "Location", "Bids as of", "Age", "Status", "Rows"]
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(row=r, column=j, value=h)
        cell.font = ST.F_HEADER
        cell.fill = ST.FILL_HEADER
    r += 1
    for h in health:
        ws.cell(row=r, column=1, value=h.company).font = ST.F_BODY
        ws.cell(row=r, column=2, value=h.location).font = ST.F_BODY
        c = ws.cell(row=r, column=3, value=h.as_of)
        c.number_format = ST.FMT_DATE
        c.font = ST.F_BODY
        ws.cell(row=r, column=4, value=h.age_days if h.age_days is not None else "").font = ST.F_BODY
        sc = ws.cell(row=r, column=5, value=h.status)
        sc.fill, sc.font = ST.badge(h.status)
        ws.cell(row=r, column=6, value=h.rows).font = ST.F_BODY
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="Best cash bids").font = ST.F_H2
    r += 1
    heads = ["Commodity", "When", "Best Bid", "Basis", "Delivery", "Company", "Location"]
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(row=r, column=j, value=h)
        cell.font = ST.F_HEADER
        cell.fill = ST.FILL_HEADER
    r += 1
    dec = eff.get("decimals", {"cash": 4, "basis": 4})
    for b in best:
        ws.cell(row=r, column=1, value=b.commodity).font = ST.F_BODY
        ws.cell(row=r, column=2, value=b.window).font = ST.F_BODY
        c = ws.cell(row=r, column=3, value=b.cash_price)
        c.number_format = ST.fmt_cash(dec["cash"])
        c.font = ST.F_BODY_BOLD
        c.fill = ST.FILL_BEST
        cb = ws.cell(row=r, column=4, value=b.basis)
        cb.number_format = ST.fmt_basis(dec["basis"])
        cb.font = ST.F_BODY
        ws.cell(row=r, column=5, value=b.delivery_label).font = ST.F_BODY
        ws.cell(row=r, column=6, value=b.company).font = ST.F_BODY
        ws.cell(row=r, column=7, value=b.location).font = ST.F_BODY
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="Needs attention").font = ST.F_H2
    r += 1
    if not needs_attention:
        c = ws.cell(row=r, column=1, value="Everything looks up to date.")
        c.font = ST.F_BODY_OK
    else:
        for item in needs_attention:
            bc = ws.cell(row=r, column=1, value=item["badge"])
            bc.fill, bc.font = ST.badge(item["badge"])
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            mc = ws.cell(row=r, column=2, value=item["message"])
            mc.font = ST.F_BODY
            mc.alignment = ST.A_WRAP
            r += 1


def write_data_health(ws, health: list, eff, now):
    ws.sheet_view.showGridLines = False
    widths = [22, 16, 30, 12, 8, 10, 14, 40, 46]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _title_block(ws, 9, "Data Health",
                 "How fresh is each source page, and anything worth checking",
                 f"Generated {now:%m/%d/%Y %I:%M %p}", None, eff)
    header_row = 6
    heads = ["Company", "Location", "Page", "Bids as of", "Age", "Status",
             "Same as last?", "What we noticed", "Suggested action"]
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(row=header_row, column=j, value=h)
        cell.font = ST.F_HEADER
        cell.fill = ST.FILL_HEADER
    r = header_row + 1
    for h in health:
        ws.cell(row=r, column=1, value=h.company).font = ST.F_BODY
        ws.cell(row=r, column=2, value=h.location).font = ST.F_BODY
        pc = ws.cell(row=r, column=3, value=(h.url or "").split("//")[-1])
        pc.font = ST.F_LINK if h.url else ST.F_BODY
        if h.url and eff.get("hyperlinks", True):
            pc.hyperlink = Hyperlink(ref=pc.coordinate, target=h.url)
        c = ws.cell(row=r, column=4, value=h.as_of)
        c.number_format = ST.FMT_DATE
        c.font = ST.F_BODY
        ws.cell(row=r, column=5, value=h.age_days if h.age_days is not None else "").font = ST.F_BODY
        sc = ws.cell(row=r, column=6, value=h.status)
        sc.fill, sc.font = ST.badge(h.status)
        ws.cell(row=r, column=7, value=h.same_as_last).font = ST.F_BODY
        nc = ws.cell(row=r, column=8, value=h.notices)
        nc.font = ST.F_BODY
        nc.alignment = ST.A_WRAP
        sug = ws.cell(row=r, column=9, value=h.suggestion)
        sug.font = ST.F_BODY
        sug.alignment = ST.A_WRAP
        r += 1
    ws.freeze_panes = f"A{header_row + 1}"
    if r > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:I{r - 1}"

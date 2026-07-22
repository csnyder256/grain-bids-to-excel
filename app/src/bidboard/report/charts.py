"""Native openpyxl charts, fed from a hidden _ChartData sheet so data sheets
stay clean and every reference stays inside the workbook."""
from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, Reference, Series

from . import styles

CHART_DATA_SHEET = "_ChartData"


def ensure_chart_data_sheet(wb):
    if CHART_DATA_SHEET in wb.sheetnames:
        return wb[CHART_DATA_SHEET]
    ws = wb.create_sheet(CHART_DATA_SHEET)
    ws.sheet_state = "hidden"
    ws["A1"] = "chart data (do not edit)"
    return ws


def _write_block(data_ws, title: str, categories: list[str],
                 series: dict[str, list[float]]) -> tuple[int, int, int]:
    """Write a labeled block; return (header_row, first_data_row, last_data_row)."""
    start = (data_ws.max_row or 1) + 2
    data_ws.cell(row=start, column=1, value=title)
    header_row = start + 1
    data_ws.cell(row=header_row, column=1, value="Category")
    for j, name in enumerate(series, start=2):
        data_ws.cell(row=header_row, column=j, value=name)
    for i, cat in enumerate(categories):
        r = header_row + 1 + i
        data_ws.cell(row=r, column=1, value=cat)
        for j, (name, values) in enumerate(series.items(), start=2):
            data_ws.cell(row=r, column=j, value=values[i])
    return header_row, header_row + 1, header_row + len(categories)


def add_bid_bar_chart(wb, sheet, data_ws, title, categories, values, anchor) -> bool:
    if len(categories) < 2:
        return False
    hr, first, last = _write_block(data_ws, f"{title} (bars)", categories,
                                   {"Cash bid": values})
    chart = BarChart()
    chart.type = "col"
    chart.gapWidth = 60
    chart.title = title
    chart.legend = None
    chart.height = 8
    chart.width = 16
    data = Reference(data_ws, min_col=2, min_row=hr, max_row=last)
    cats = Reference(data_ws, min_col=1, min_row=first, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.numFmt = '"$"0.00'
    chart.y_axis.majorGridlines = None
    sheet.add_chart(chart, anchor)
    return True


def add_basis_line_chart(wb, sheet, data_ws, title, categories, values, anchor) -> bool:
    if len(categories) < 3:
        return False
    hr, first, last = _write_block(data_ws, f"{title} (basis)", categories,
                                   {"Basis": values})
    chart = LineChart()
    chart.title = title
    chart.legend = None
    chart.height = 8
    chart.width = 16
    chart.style = 2
    data = Reference(data_ws, min_col=2, min_row=hr, max_row=last)
    cats = Reference(data_ws, min_col=1, min_row=first, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.numFmt = "0.0000"
    sheet.add_chart(chart, anchor)
    return True

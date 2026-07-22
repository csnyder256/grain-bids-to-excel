"""Report orchestration. build_reports() is the pure entry (inject clock +
paths for tests); build_for_run() is the app-facing convenience the
ScanManager and Results page call."""
from __future__ import annotations

import logging
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from .. import settings as settings_mod
from . import files as F
from . import health as HEALTH
from . import layout as L
from . import queries as Q
from .naming import sheet_name

log = logging.getLogger("bidboard.report.builder")


@dataclass
class ReportFile:
    path: str
    mode: str
    scope: str
    sheets: int
    rows: int


@dataclass
class BuildResult:
    build_id: int | None
    files: list[ReportFile] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    primary: str | None = None


def build_reports(
    store, settings: dict, run_id: int, out_dir: Path,
    *, now: datetime, record_build: bool = True,
) -> BuildResult:
    result = BuildResult(build_id=None)
    modes = settings["output"].get("workbook_modes", ["master"])
    tokens_base = {"date": now.strftime("%Y-%m-%d"), "time": now.strftime("%H%M"),
                   "scan_id": str(run_id)}

    for mode in modes:
        if mode == "master":
            result.files.extend(_build_master(store, settings, run_id, out_dir, now, tokens_base))
        elif mode == "per_company":
            result.files.extend(_build_per_company(store, settings, run_id, out_dir, now, tokens_base))
        elif mode == "per_commodity":
            result.files.extend(_build_per_commodity(store, settings, run_id, out_dir, now, tokens_base))

    # archive old files
    if settings["output"].get("archive", {}).get("enabled", True):
        keep = settings["output"]["archive"].get("keep_latest", 10)
        result.warnings.extend(F.archive_old(out_dir, keep, now.strftime("%Y-%m")))

    if result.files:
        result.primary = result.files[0].path

    if record_build:
        result.build_id = store.record_build(
            run_id, settings, [asdict(f) for f in result.files]
        )
    return result


def _new_wb() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _all_rows(store, run_id, settings):
    eff = settings_mod.resolve(settings, None)
    return Q.fetch_bid_rows(store, run_id, eff)


def _build_master(store, settings, run_id, out_dir, now, tokens):
    eff_global = settings_mod.resolve(settings, None)
    companies = Q.distinct_companies_in_run(store, run_id)
    all_rows = Q.fetch_bid_rows(store, run_id, eff_global)
    charts = settings["output"].get("charts", "none")

    wb = _new_wb()
    taken: set[str] = set()

    # Summary cover
    cover = wb.create_sheet(sheet_name("Summary", taken))
    health = HEALTH.compute_health(store, run_id, settings["output"]["stale_thresholds"], now)
    best = HEALTH.best_bids(all_rows, now)
    attention = _attention_items(health)
    meta = {"company_count": len(companies), "bid_count": len(all_rows)}
    L.write_cover(cover, meta, health, best, attention, eff_global, now)

    # Data Health
    if settings["output"].get("include_health_sheet", True):
        hs = wb.create_sheet(sheet_name("Data Health", taken))
        L.write_data_health(hs, health, eff_global, now)

    # one sheet per company
    sheet_per = settings["output"].get("master", {}).get("sheet_per", "company")
    data_sheets = 0
    cap = settings["output"].get("max_data_sheets_per_workbook", 24)
    for comp in companies:
        eff = settings_mod.resolve(settings, comp["id"])
        rows = Q.fetch_bid_rows(store, run_id, eff, company_id=comp["id"])
        if not rows:
            continue
        if data_sheets >= cap:
            log.info("sheet cap reached at %d", cap)
            break
        ws = wb.create_sheet(sheet_name(comp["name"], taken))
        as_of = _max_as_of(rows)
        L.write_bid_sheet(
            ws, rows, eff,
            title=comp["name"],
            subtitle=f"Cash bids as of {as_of:%B %d, %Y}" if as_of else "Cash bids",
            small=f"Scanned {now:%b %d, %Y %I:%M %p} · Source: {_domain(rows)}",
            source_url=rows[0].get("site_url"),
            include_company=False,
            include_location=(sheet_per == "company"),
            charts_mode=eff.get("charts", charts),
            now=now,
        )
        data_sheets += 1

    # All Bids flat sheet
    if settings["output"].get("master", {}).get("include_all_bids_sheet", True) and all_rows:
        ws = wb.create_sheet(sheet_name("All Bids", taken))
        L.write_bid_sheet(
            ws, all_rows, eff_global,
            title="All Bids",
            subtitle=f"Every bid collected in this scan · {len(all_rows)} rows",
            small=f"Scanned {now:%b %d, %Y %I:%M %p}",
            source_url=None, include_company=True, include_location=True,
            charts_mode="none", now=now,
        )

    if not wb.sheetnames:
        wb.create_sheet("Summary")
    pattern = settings["output"]["filename_patterns"]["master"]
    filename = F.render_filename(pattern, tokens)
    target = F.unique_path(out_dir, filename)
    written = F.safe_write(wb, target)
    return [ReportFile(str(written), "master", "All",
                       len(wb.sheetnames), len(all_rows))]


def _build_per_company(store, settings, run_id, out_dir, now, tokens):
    out = []
    for comp in Q.distinct_companies_in_run(store, run_id):
        eff = settings_mod.resolve(settings, comp["id"])
        rows = Q.fetch_bid_rows(store, run_id, eff, company_id=comp["id"])
        if not rows:
            continue
        wb = _new_wb()
        taken: set[str] = set()
        as_of = _max_as_of(rows)
        ws = wb.create_sheet(sheet_name(comp["name"], taken))
        L.write_bid_sheet(
            ws, rows, eff, title=comp["name"],
            subtitle=f"Cash bids as of {as_of:%B %d, %Y}" if as_of else "Cash bids",
            small=f"Scanned {now:%b %d, %Y %I:%M %p} · Source: {_domain(rows)}",
            source_url=rows[0].get("site_url"),
            include_company=False, include_location=True,
            charts_mode=eff.get("charts", "none"), now=now,
        )
        t = dict(tokens, company=comp["name"])
        filename = F.render_filename(settings["output"]["filename_patterns"]["per_company"], t)
        written = F.safe_write(wb, F.unique_path(out_dir, filename))
        out.append(ReportFile(str(written), "per_company", comp["name"],
                              len(wb.sheetnames), len(rows)))
    return out


def _build_per_commodity(store, settings, run_id, out_dir, now, tokens):
    out = []
    eff_global = settings_mod.resolve(settings, None)
    for commodity in Q.distinct_commodities_in_run(store, run_id):
        rows = Q.fetch_bid_rows(store, run_id, dict(eff_global, commodities=[commodity]))
        rows = [r for r in rows if r.get("commodity") == commodity]
        if not rows:
            continue
        wb = _new_wb()
        taken: set[str] = set()
        # one sheet per company for this commodity
        companies = sorted({r["company_name"] for r in rows})
        for cname in companies:
            crows = [r for r in rows if r["company_name"] == cname]
            ws = wb.create_sheet(sheet_name(cname, taken))
            L.write_bid_sheet(
                ws, crows, dict(eff_global, group_by="location"),
                title=f"{cname} - {commodity}",
                subtitle=f"{commodity} cash bids", small=f"Scanned {now:%b %d, %Y}",
                source_url=crows[0].get("site_url"),
                include_company=False, include_location=True,
                charts_mode="none", now=now,
            )
        t = dict(tokens, commodity=commodity)
        filename = F.render_filename(settings["output"]["filename_patterns"]["per_commodity"], t)
        written = F.safe_write(wb, F.unique_path(out_dir, filename))
        out.append(ReportFile(str(written), "per_commodity", commodity,
                              len(wb.sheetnames), len(rows)))
    return out


def _attention_items(health):
    items = []
    for h in health:
        if h.status in ("STALE", "FAILED"):
            items.append({"badge": "STALE" if h.status == "STALE" else "FAILED",
                          "message": f"{h.company} – {h.location}: {h.notices}"})
        elif "noticed" in (h.notices or "").lower() or h.suggestion:
            if h.status != "FRESH":
                items.append({"badge": "CHECK",
                              "message": f"{h.company} – {h.location}: {h.notices}"})
    return items


def _max_as_of(rows):
    dates = []
    for r in rows:
        v = r.get("page_as_of")
        if v:
            try:
                from datetime import date
                dates.append(date.fromisoformat(str(v)[:10]))
            except ValueError:
                pass
    return max(dates) if dates else None


def _domain(rows):
    url = rows[0].get("site_url") if rows else ""
    return (url or "").split("//")[-1].split("/")[0]


# ---------------- app-facing convenience ----------------

def build_for_run(store, settings_path, run_id: int) -> dict | None:
    """Called by the ScanManager after a scan and by the Results 'Rebuild'
    button. Returns a small dict for the UI, or None on nothing-to-build."""
    from ..paths import SPREADSHEETS_DIR

    settings = settings_mod.load_settings(settings_path)
    result = build_reports(
        store, settings, run_id, SPREADSHEETS_DIR, now=datetime.now()
    )
    if not result.files:
        return None
    return {
        "build_id": result.build_id,
        "primary": result.primary,
        "files": [f.path for f in result.files],
        "count": len(result.files),
        "warnings": result.warnings,
    }

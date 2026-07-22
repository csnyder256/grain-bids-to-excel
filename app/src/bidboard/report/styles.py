"""The single theme module - every Font/Fill/Border/format for the
workbooks lives here so 30 sheets stay visually identical and a rebrand is
a one-file change. Calibri is guaranteed present on Excel 2016+."""
from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---- palette (hex, no leading #) ----
BRAND_DARK = "1F3B2C"    # deep pine - titles, header fill
BRAND_MID = "3E6B4F"     # mid green - links, chart series 1
BRAND_ACCENT = "C8A44D"  # harvest gold - divider rules, chart series 2
INK = "212B26"
MUTED = "6B7671"
BAND = "F2F5F3"
GRID = "D5DBD8"
POS_GREEN, POS_FILL = "1E7B34", "E7F4EA"
NEG_RED, NEG_FILL = "B3261E", "FBEAE9"
FRESH_FILL, AGING_FILL, STALE_FILL = "D1E7DD", "FFF3CD", "F8D7DA"
FRESH_TEXT, AGING_TEXT, STALE_TEXT = "0F5132", "664D03", "842029"
ACTION_FILL, ACTION_TEXT = "CFE2FF", "084298"
CASH_SCALE_TOP = "CDE3D3"
WHITE = "FFFFFF"

# ---- fonts ----
F_TITLE = Font(name="Calibri Light", size=18, bold=True, color=BRAND_DARK)
F_SUBTITLE = Font(name="Calibri", size=11, italic=True, color=MUTED)
F_H2 = Font(name="Calibri", size=13, bold=True, color=BRAND_DARK)
F_HEADER = Font(name="Calibri", size=11, bold=True, color=WHITE)
F_BODY = Font(name="Calibri", size=11, color=INK)
F_BODY_BOLD = Font(name="Calibri", size=11, bold=True, color=INK)
F_SMALL = Font(name="Calibri", size=9, italic=True, color=MUTED)
F_LINK = Font(name="Calibri", size=11, color=BRAND_MID, underline="single")

# ---- fills ----
FILL_HEADER = PatternFill("solid", fgColor=BRAND_DARK)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_RULE = PatternFill("solid", fgColor=BRAND_ACCENT)
FILL_BEST = PatternFill("solid", fgColor=CASH_SCALE_TOP)

# ---- borders ----
S_HAIR = Side(style="hair", color=GRID)
S_GOLD = Side(style="thin", color=BRAND_ACCENT)
S_SECTION = Side(style="medium", color=BRAND_DARK)
B_ROW = Border(bottom=S_HAIR)
B_HEADER = Border(bottom=S_GOLD)
B_GROUPTOP = Border(top=S_SECTION, bottom=S_HAIR)

# ---- alignments ----
A_LEFT = Alignment(horizontal="left", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")
A_CENTER = Alignment(horizontal="center", vertical="center")
A_TITLE = Alignment(horizontal="left", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


# ---- number formats (regenerated from the decimals setting) ----
def fmt_cash(d: int = 4) -> str:
    return '"$"#,##0.' + "0" * d


def fmt_basis(d: int = 4) -> str:
    z = "0." + "0" * d
    return f"+{z};-{z};{z}"


def fmt_change(d: int = 4) -> str:
    z = "0." + "0" * d
    return f'+{z};-{z};"unch"'


def fmt_futures(d: int = 4) -> str:
    return "#,##0." + "0" * d


FMT_DATE = "mm/dd/yyyy"
FMT_DATETIME = "mm/dd/yyyy h:mm AM/PM"
FMT_AGE = '0 "days"'


# ---- conditional-format rule factories ----
def cf_pos_neg(rng: str) -> list:
    return [
        CellIsRule(operator="greaterThan", formula=["0"],
                   font=Font(color=POS_GREEN),
                   fill=PatternFill("solid", fgColor=POS_FILL)),
        CellIsRule(operator="lessThan", formula=["0"],
                   font=Font(color=NEG_RED),
                   fill=PatternFill("solid", fgColor=NEG_FILL)),
    ]


def cf_cash_scale() -> ColorScaleRule:
    return ColorScaleRule(
        start_type="min", start_color=WHITE,
        end_type="max", end_color=CASH_SCALE_TOP,
    )


BADGE_FILLS = {
    "FRESH": (FRESH_FILL, FRESH_TEXT),
    "AGING": (AGING_FILL, AGING_TEXT),
    "STALE": (STALE_FILL, STALE_TEXT),
    "FAILED": (STALE_FILL, STALE_TEXT),
    "ACTION": (ACTION_FILL, ACTION_TEXT),
    "CHECK": (AGING_FILL, AGING_TEXT),
}
